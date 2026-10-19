"""Genera la guía de operación y defensa; ejecutar con Python y python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
header = section.header.paragraphs[0]
header.text = 'Sky Import | Operación y exposición'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Página ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, level=2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = text
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DDEAF2' if ri == 0 else ('F3F5F6' if ri % 2 == 0 else 'FFFFFF')); tcPr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:'+edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
            tcPr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:'+side); el.set(qn('w:w'), '110'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcPr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs: run.font.size = Pt(9); run.bold = ri == 0
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    p('')

doc.add_heading('Guía de operación y exposición de Sky Import', 0)
p('Revisión del 19 de octubre de 2026 · Preparación de la entrega final de diciembre')
p('Esta guía sirve para preparar el entorno, operar la tienda y demostrar el proyecto. Acompaña al SRS versión 2.3 y conserva como antecedente la versión 2.2 presentada en papel para aval.')
h('Qué resuelve el proyecto')
p('Sky Import permite consultar componentes para PC, comprobar reglas de compatibilidad y preparar un pedido en español o portugués. Los precios se almacenan en dólares y se presentan también en guaraníes y reales. La venta se coordina por WhatsApp; el software no cobra ni emite facturas fiscales.')
h('Los tres recorridos que hay que distinguir')
table(['Recorrido', 'Resultado'], [
    ['Compra directa', 'Prepara el mensaje de una pieza. No crea un pedido ni descuenta stock.'],
    ['Checkout registrado', 'Revalida precios y existencias, registra el pedido y reserva unidades antes de abrir WhatsApp.'],
    ['Continuación por fallo técnico', 'Puede preparar un mensaje sin registro, sujeto a confirmación comercial. Un rechazo de disponibilidad bloquea el pedido.'],
], [5, 12])
h('Estado de los entornos')
p('Las correcciones se publicaron el 19/10/2026 en https://sky-import-jet.vercel.app. La base tiene aplicada la migración de integridad de pedidos y cupones. Antes del cambio se restauró una copia lógica en una base local aislada; el ciclo de reserva y cancelación se comprobó también en producción con reversión de la prueba.')
p('Abrir WhatsApp no demuestra que se haya enviado el mensaje, confirmado la venta o recibido el pago. Esta distinción debe mantenerse en la exposición y en la operación del comercio.')

page('Preparar el entorno de ensayo')
h('Instalación reproducible')
p('Usar Node.js 24.19.0, la versión fijada en .nvmrc, y las dependencias de package-lock.json. Desde la carpeta del proyecto, ejecutar:')
for line in ['npm ci', 'npm run verify', 'npm start']:
    r=p(line).runs[0]; r.font.name='Consolas'; r.font.size=Pt(10)
p('La compilación genera las páginas de producción. npm start las sirve en http://localhost:3000. Para editar durante el desarrollo se utiliza npm run dev. Las pruebas de navegador levantan su propio servidor en el puerto 3100.')
h('Conexión a los datos')
p('Las variables están en .env.local y su plantilla es .env.example. La URL y la clave pública identifican Supabase. La clave de servicio permanece en el servidor y permite la gestión autorizada de cuentas; no debe mostrarse en capturas ni proyectarse.')
p('La web local puede apuntar a la misma base remota que el dominio. Ejecutarla en localhost no aísla los datos. Para ensayar altas, cambios de precios, cancelaciones o borrados, utilizar un entorno de prueba o registros expresamente preparados para ese ensayo.')
h('Preparación antes de exponer')
p('Comprobar conexión, abrir catálogo y panel, validar el usuario y seleccionar los productos que se usarán. Preparar una combinación compatible y otra incompatible en el configurador. Confirmar que la moneda y el idioma estén en el estado previsto.')
p('Tener a mano el SRS revisado, esta guía y las capturas de los recorridos verificados. Si falla internet, puede mostrarse el catálogo de respaldo y explicarse su límite: no acredita operaciones del panel ni un registro exitoso en la base.')
h('Comprobación automatizada')
p('npm run verify revisa lint, tipos, pruebas unitarias y compilación. npm run e2e comprueba los recorridos de navegador. npm run size verifica el presupuesto de JavaScript. Las cuentas temporales requieren RUN_ACCOUNT_E2E=1; se eliminan al terminar.')
p('Las transacciones se comprueban con node tests/database/pedidos.mjs sobre PostgreSQL aislado. La preparación de sus dependencias está indicada al inicio del script y no modifica la base comercial.')

page('Operar el panel de administración')
table(['Área', 'Operación y comprobación'], [
    ['Productos y variantes', 'Crear o editar la pieza. Guardar y volver a abrir la ficha para verificar los valores. Desactivar la visibilidad retira el producto de la tienda tras revalidar.'],
    ['Existencias', 'Sin variantes activas se edita el stock base. Con variantes activas, editar cada variante; catálogo, listado y alertas suman sus unidades.'],
    ['Precios', 'Introducir la tasa como unidades por dólar, sin separadores de miles. Guardar y comprobar un precio convertido en la tienda.'],
    ['Cupones', 'Definir tipo, valor, compra mínima, vigencia, límite y habilitación. Probar un código válido y otro rechazado.'],
    ['Pedidos', 'Consultar detalle e importes históricos. Cambiar estado o agregar una nota. Cancelar libera la reserva una vez; reabrir vuelve a exigir existencias.'],
    ['Configuración', 'Cada usuario cambia su contraseña. El moderador puede crear administradores, restablecer claves ajenas y retirar o restituir accesos.'],
], [4, 13])
h('Estados y permisos')
p('Administrador: gestiona inventario, precios, cupones y seguimiento. Moderador: añade la gestión de cuentas y la baja de pedidos. Los permisos se comprueban en el servidor y en las políticas de la base; ocultar un botón no constituye la protección.')
p('Un pedido entregado no admite cambios de estado. Eliminarlo no repone existencias. Eliminar un pedido cancelado tampoco repite la restitución. El borrado exige confirmar su número y elimina sus líneas; para una cancelación comercial ordinaria se conserva el registro usando el estado Cancelado.')
h('Interpretar el tablero')
p('El tablero utiliza hasta 50 pedidos recientes y el historial hasta 200. El indicador de pedidos no cancelados no acredita cobros y su conversión agregada usa la tasa vigente. El detalle de cada pedido conserva sus valores históricos.')
p('Un aviso de error de consulta no significa que la base esté vacía. Reintentar y comprobar la conexión antes de importar de nuevo el catálogo o modificar datos.')

page('Guion de demostración')
p('Duración sugerida de ensayo: 12 minutos. Ajustar al tiempo que indique el docente. Preparar de antemano las pestañas y los datos de prueba; no improvisar cambios sobre registros comerciales.')
table(['Tiempo', 'Mostrar', 'Explicar y comprobar'], [
    ['0–1 min', 'Problema y alcance', 'Catálogo, compatibilidad y preparación del pedido. El pago y la logística quedan fuera del software.'],
    ['1–3 min', 'Catálogo y ficha', 'Buscar, filtrar y cambiar idioma y moneda. Mostrar código, especificaciones, disponibilidad y variante cuando corresponda. RF02 a RF04 y RF08.'],
    ['3–5 min', 'Arma tu PC', 'Mostrar una combinación compatible y una incompatible. Explicar la regla concreta y distinguir bloqueo de advertencia. RF07.'],
    ['5–7 min', 'Carrito y checkout', 'Cambiar cantidad, recargar, introducir datos de prueba y aplicar un cupón. Revisar el mensaje preparado sin enviarlo al comercio. RF05, RF06 y RF09 a RF11.'],
    ['7–10 min', 'Panel', 'Mostrar inventario, alertas, tasa y pedido registrado. En entorno de ensayo, cancelar y comprobar que las unidades vuelven una sola vez. RF12 a RF16 y RF20.'],
    ['10–11 min', 'Usuarios y seguridad', 'Diferenciar administrador y moderador. Mostrar que las funciones de cuentas no están disponibles al administrador. RF01 y RF17 a RF21.'],
    ['11–12 min', 'Evidencia y cierre', 'Relacionar una acción con su requisito, diagrama y prueba. Exponer los límites verificados y las mejoras futuras.'],
], [2.1, 4.2, 10.7])
h('Apertura sugerida')
p('“Sky Import organiza la consulta de componentes y la preparación del pedido de un comercio de Ciudad del Este. El comprador puede comparar especificaciones, comprobar compatibilidad y ver precios en tres monedas. El administrador mantiene los datos desde un panel, y la atención comercial confirma la venta por WhatsApp.”')
h('Ejemplo para explicar la integridad')
p('“El navegador envía qué producto quiere el cliente y cuántas unidades. La base vuelve a leer el precio, valida el stock y registra todo en una transacción. Si una condición falla, el pedido no queda guardado a medias.”')

page('Preparar las respuestas de la defensa')
h('Por qué se eligió esta arquitectura')
p('Next.js permite generar las páginas públicas y ejecutar en el servidor las acciones del panel. PostgreSQL conserva las reglas transaccionales junto a los datos. Supabase aporta autenticación y acceso a PostgreSQL; las políticas RLS limitan las operaciones según la cuenta.')
h('Qué ocurre si alguien cambia el precio desde el navegador')
p('El registro del pedido no acepta el precio del cliente como fuente de verdad. Relee los importes de la base y calcula el descuento dentro de la transacción. El mensaje de un pedido registrado usa los importes devueltos por esa operación.')
h('Cómo se evitan reservas duplicadas o cupones excedidos')
p('Las filas afectadas se bloquean durante el registro. El cupón se revalida bajo bloqueo antes de incrementar sus usos. Al cancelar se marca la reserva como liberada; al borrar no se devuelve por segunda vez. Reabrir exige existencias nuevamente.')
h('Qué significa tolerancia a fallos')
p('Un fallo técnico puede permitir continuar por WhatsApp sin guardar un pedido. Esto conserva el contacto comercial, pero no garantiza un número, stock reservado ni seguimiento interno. Un rechazo comercial por disponibilidad se trata de otra manera y bloquea la compra.')
h('Qué demuestran las pruebas')
p('Las unitarias comprueban reglas y transformaciones. Las de navegador comprueban interacción y accesibilidad automatizada. PostgreSQL aislado comprueba transacciones, concurrencia y políticas. Las pruebas interceptadas del checkout no prueban por sí solas la base publicada. Tampoco una auditoría automática acredita toda la accesibilidad.')
h('Qué sigue fuera del alcance')
p('Cobro electrónico, factura fiscal, seguimiento de transporte, cuentas de compradores y bitácora completa de operaciones. Los tiempos de respuesta y la disponibilidad sostenida requieren mediciones específicas. La restauración de datos está probada localmente; faltan copias diarias y externas y una recuperación integral del servicio.')

page('Cerrar la entrega operativa')
h('Publicación de la revisión')
p('1. Publicación comprobada el 19/10/2026: sky-import-jet.vercel.app, en Vercel de Ciendel, conectado al proyecto Supabase bhlkdinovdhtwxkkujic.')
p('2. La migración 20261019141606_integridad_pedidos_cupones.sql está aplicada y registrada. Antes se guardó y restauró una copia lógica. No volver a ejecutar versiones antiguas de funciones.')
p('3. Verificación posterior: 10 pruebas de navegador correctas entre escritorio y móvil; panel, permisos, accesibilidad y URL canónicas. Los rechazos comerciales del checkout se simularon; el ciclo de stock se probó por SQL en producción con reversión.')
p('4. Resultado final de los datos: 37 productos, cero pedidos y cero cuentas temporales de prueba. En futuros cambios, conservar fecha, versión y resultados y repetir el recorrido afectado.')
h('Respaldo y restauración')
p('Ejecución del 19/10/2026: copia privada previa a la migración, fuera del repositorio, con SQL de reconstrucción, informe y sumas SHA-256. Se restauraron nueve tablas públicas y los registros de dos tablas de Auth; todos coincidieron. También se reconstruyeron 17 políticas RLS y se aplicó la migración sobre la copia.')
p('Alcance: no incluye sesiones, tokens de renovación, MFA, configuración de proveedores ni secretos del despliegue. Storage estaba vacío. La prueba no reconstruye todo el servicio Supabase. El plan Free no ofrece copias administradas: el titular debe conservar una copia externa privada y organizar la frecuencia diaria. La ubicación y el procedimiento están en README; nunca adjuntar los datos de Auth al documento académico.')
h('Antes de imprimir y exponer')
p('Revisar índice, numeración, figuras y fechas del SRS. Confirmar que la versión demostrada corresponde al documento. Ensayar el recorrido completo y repetir las comprobaciones después de cualquier cambio. Conservar la versión impresa de octubre como antecedente, sin sustituir sus resultados históricos.')
p('La entrega final queda cerrada cuando la migración y el despliegue estén comprobados, el recorrido de demostración pueda repetirse y los pendientes operativos tengan evidencia o un alcance acordado con el docente.')

out = ROOT / 'docs/presentacion/Guia-Operacion-y-Exposicion.docx'
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
